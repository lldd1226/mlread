from pathlib import Path
from bs4 import BeautifulSoup, Tag, NavigableString
from typing import List, Optional, Dict, Tuple
from dataclasses import dataclass, field
import re
import MEGAok
import openpyxl

# ── 全局配置 ─────────────────────────────────────────────────────────────────

CSS_HREF:         str = "../../mega.css"


# ── IR dataclasses（解析层输出，重组层的唯一输入）────────────────────────────

@dataclass
class RawFootnoteGroup:
    label:      Optional[str]   # None = 本页开头无标签，是上页脚注的跨页续行
    paragraphs: List[str]       # 已渲染的 HTML 段落列表


@dataclass
class Block:
    type:      str              # 'p' | 'h1'–'h6' | 'editorial' | 'note' | 'table'
    html:      str  = ''        # 已渲染的内联 HTML（note 类型为空）
    level:     int  = 0         # 标题原始级别 1–8；其他块为 0
    indent:    bool = False     # 首行缩进
    quote:     bool = False     # 块级引用段落
    footnotes: List[RawFootnoteGroup] = field(default_factory=list)


@dataclass
class Footnote:
    old_num:      str
    new_num:      str
    page_num:     int
    content_html: str
    extra_paras:  List[str] = field(default_factory=list)
class FootnoteNumberer:
    def __init__(self) -> None:
        self._assigned:   Dict[str, int]             = {}
        self._fn_map:     Dict[Tuple[int, str], str] = {}
        self._ref_counts: Dict[str, int]             = {}

    @property
    def fn_map(self) -> Dict[Tuple[int, str], str]:
        return self._fn_map

    def assign(self, page_num: int, old_num: str) -> Optional[str]:
        if not old_num:
            return None
        key = (page_num, old_num)
        if key in self._fn_map:
            return self._fn_map[key]
        cnt = self._assigned.get(old_num, 0) + 1
        self._assigned[old_num] = cnt
        new_num = old_num if cnt == 1 else f"{old_num}-{cnt}"
        self._fn_map[key] = new_num
        return new_num

    def next_ref_anchor(self, new_num: str) -> Tuple[str, str]:
        cnt = self._ref_counts.get(new_num, 0) + 1
        self._ref_counts[new_num] = cnt
        anchor_id = f"ZM{new_num}" if cnt == 1 else f"ZM{new_num}-{cnt}"
        return anchor_id, f"#M{new_num}"

    def ref_anchors_for(self, new_num: str) -> List[str]:
        cnt = self._ref_counts.get(new_num, 0)
        if cnt == 0:
            return [f"ZM{new_num}"]
        result = [f"ZM{new_num}"]
        result.extend(f"ZM{new_num}-{i}" for i in range(2, cnt + 1))
        return result


class MegaEtxMerger:
    _PLACEHOLDER_PAT = re.compile(r'\x00(FN[RA])([\da-z\[\]]+)\x00')

    def __init__(self, volume: int, page_group,
                 input_dir: Path = None, output_dir: Path = None,volume_info:str=""):
        self.volume        = volume
        self.raw_dir       = input_dir  or Path(f"II_{volume}/raw")
        self.out_dir       = output_dir or Path(f"II_{volume}/merged")
        self.use_footnotes = [5,10,12,13,15]
        self.cutcontent = [5,10,13,15]
        self.volume_info=volume_info
        if page_group and isinstance(page_group[0], (list, tuple)):
            self.groups = [list(g) for g in page_group if g]
        else:
            self.groups = [list(page_group)]

    def _path(self, page: int) -> Path:
        return self.raw_dir / f"p{page:04d}.html"

    # ═══════════════════════════════════════════════════════════════════════════
    # 解析层：soup → IR（Block / RawFootnoteGroup）
    # ═══════════════════════════════════════════════════════════════════════════

    def _inline(self, tag: Tag) -> str:
        buf = []
        for ch in tag.children:
            if isinstance(ch, NavigableString):
                buf.append(re.sub(r'[\s\r\n]{6,}', ' ', str(ch)))
                continue
            if not isinstance(ch, Tag):
                continue
            cls = ch.get('class', [])
            
            # --- 脚注引用 ---
            if ch.name == 'sup' and 'megaEtx' in cls:
                num = re.sub(r'[^0-9a-z\]\[]+', '', ch.get_text())
                num_match = re.match(r"^([\d]+[\]\[a-z]*)[\)]*", num)
                nxt = ch.next_sibling
                if isinstance(nxt, NavigableString) and re.match(r'\s*\)', str(nxt)):
                    nxt.replace_with(NavigableString(
                        re.sub(r'^\s*\)', '', str(nxt))))
                buf.append(f'\x00FNR{num_match.group(1)}\x00' if num_match
                           else f'<sup>{self._inline(ch)}</sup>')
            
            # --- 公式（绝对不能让 BS4 破坏 MathML 结构，原样输出并剔除 megaEtx）---
            elif 'formula' in cls:
                buf.append(re.sub(r'\s*class="[^"]*megaEtx[^"]*"', '', str(ch)))
            
            # --- 上下标 ---
            elif ch.name in ('sup', 'sub'):
                buf.append(f'<{ch.name}>{self._inline(ch)}</{ch.name}>')
            elif ch.name == 'br':
                pass
            
            # --- 文本格式化 ---
            elif 'em' in cls or 'hi' in cls:  # 新增 'hi' 支持 (如 <span class="hi i">)
                buf.append(f'<i>{self._inline(ch)}</i>')
            elif 'sp' in cls:
                buf.append(f'<span class="sp">{self._inline(ch)}</span>')
            elif 'u' in cls:
                buf.append(f'<u>{self._inline(ch)}</u>')
            elif 'dot' in cls:
                buf.append(f'<u class="dot">{self._inline(ch)}</u>')
            elif 'mpb' in cls:
                # 区分 label-mpb 和单独的 mpb，防止竖线重复或被吞
                if 'label-mpb' in cls:
                    buf.append(f'<span class="mpb">{ch.get_text(strip=True)}</span>')
                else:
                    buf.append(f'<span class="mpb">|{ch.get_text(strip=True)}|</span>')
            elif 'helv' in cls or 'add-helv' in cls:
                buf.append(ch.get_text())

            elif 'fna' in cls:
                num = re.sub(r'[^0-9a-z]+', '', ch.get_text())
                if num:
                    buf.append(f'\x00FNA{num}\x00')
            else:
                buf.append(self._inline(ch))
        return ''.join(buf)

    def _join_lines(self, line_divs: List[Tag]) -> str:
        parts = []
        for ld in line_divs:
            lc = ld.find('div', class_=re.compile(r'\blineContent\b'))
            if not lc:
                continue
            rendered = self._inline(lc).strip()
            if not rendered:
                continue
            if parts:
                if parts[-1].endswith('-'):
                    parts[-1] = parts[-1][:-1]
                    parts.append(rendered)
                else:
                    parts.append(' ')
                    parts.append(rendered)
            else:
                parts.append(rendered)
        return ''.join(parts)

    @staticmethod
    def _line_divs(block: Tag) -> List[Tag]:
        return block.find_all(
            'div',
            class_=lambda c: c and 'megaEtx' in c and 'line' in c
                             and 'lineContent' not in c and 'lineBreaker' not in c
        )

    @staticmethod
    def _raw_tags(soup: BeautifulSoup) -> List[Tag]:
        result = []
        def walk(tag):
            if not isinstance(tag, Tag):
                return
            cls     = tag.get('class', [])
            cls_str = ' '.join(cls)
            if 'megaEtx' not in cls:
                for ch in tag.children:
                    walk(ch)
                return
            # 【新增】'table' 入选顶层块
            if ('p' in cls or 'note' in cls or 'editorialHead' in cls
                    or re.search(r'\bh[1-8][ab]?\b', cls_str)
                    or 'table' in cls):
                result.append(tag)
            else:
                for ch in tag.children:
                    walk(ch)
        walk(soup)
        return result
    def _render_head_title(self, tag: Tag) -> str:
        """把标题块渲染为带上 <br> 的多行 HTML，只对 div.line 强制换行，
           内联子元素（如 span）照常拼接，同时保留断字连字符的合并逻辑。"""
        rows = []           # 最终的行列表（每一项是一行的完整 HTML）
        cur_parts = []      # 当前拼接中的文本片段

        def flush_cur():
            """把 cur_parts 合并成一行，加入 rows"""
            if cur_parts:
                # 用空格连接所有片段，处理多余空白
                line = ' '.join(cur_parts).strip()
                # 清理多余空格（例如 '   ' -> ' '）
                line = re.sub(r'\s+', ' ', line)
                if line:
                    rows.append(line)
                cur_parts.clear()

        for child in tag.children:
            if isinstance(child, NavigableString):
                text = str(child).strip()
                if text:
                    cur_parts.append(text)
            elif isinstance(child, Tag):
                # 遇到 line div → 强制换行
                if (child.name == 'div'
                        and 'line' in child.get('class', [])
                        and 'lineContent' not in child.get('class', [])):
                    flush_cur()   # 先把之前的内联内容作为一行输出
                    # 提取 lineContent 并渲染
                    lc = child.find('div', class_=re.compile(r'\blineContent\b'))
                    if lc:
                        text = self._inline(lc).strip()
                        if text:
                            rows.append(text)
                else:
                    # 内联标签（span, sup, etc.）→ 拼入当前行
                    rendered = self._inline(child).strip()
                    if rendered:
                        cur_parts.append(rendered)
        flush_cur()

        # 处理连字符合并：如果上一行以 '-' 结尾，把它和下一行合并（不插入 <br>）
        merged = []
        for line in rows:
            if merged and merged[-1].endswith('-'):
                merged[-1] = merged[-1][:-1] + line   # 去掉末尾 '-', 直接拼接
            else:
                merged.append(line)

        return '<br>'.join(merged)
    def _parse_block(self, tag: Tag) -> Optional[Block]:
        cls     = tag.get('class', [])
        cls_str = ' '.join(cls)
        # --- 标题 ---
        m = re.search(r'h([1-8])[ab]?', cls_str)
        if m:
            level = int(m.group(1))
            html = self._render_head_title(tag)
            if not html.strip():
                html = tag.get_text(strip=True)  # 降级方案
            return Block(type=f'h{min(level, 6)}', html=html, level=level) if html else None
        # --- 编者按 ---
        if 'editorialHead' in cls:
            return Block(type='editorial',
                         html=self._join_lines(self._line_divs(tag)))
                         
        # --- 【新增】表格 ---
        if 'table' in cls:
            table_tag = tag.find('table', class_='megaEtx')
            if table_tag:
                return Block(type='table', html=self._render_table(table_tag))
            return None
            
        # --- 脚注区 ---
        if 'note' in cls:
            return Block(type='note',
                         footnotes=self._parse_note_div(tag))
                         
        # --- 正文段落 ---
        if 'p' in cls:
            lds = self._line_divs(tag)
            if not lds:
                return None
            first_lc = lds[0].find('div', class_=lambda c: c and 'lineContent' in c)
            indent   = first_lc is not None and 'indent' in (first_lc.get('class', []))
            html     = self._join_lines(lds)
            if not html.strip():
                return None
            return Block(type='p', html=html, indent=indent,
                         quote=('cE' in cls or 'cS' in cls))
        return None

    def _parse_page(self, soup: BeautifulSoup) -> List[Block]:
        return [b for tag in self._raw_tags(soup)
                if (b := self._parse_block(tag)) is not None]

    # ── note div 解析 ─────────────────────────────────────────────────────────

    def _iter_note_elements(self, note_div: Tag):
        for child in note_div.children:
            if not isinstance(child, Tag):
                continue
            cls = child.get('class', [])
            if (child.name == 'div' and 'megaEtx' in cls and 'line' in cls
                    and 'lineContent' not in cls and 'lineBreaker' not in cls):
                yield 'line', child
            elif child.name == 'div' and 'lineBreaker' in cls:
                yield 'break', child
            elif child.name == 'table':
                yield 'table', child
            elif (child.name in ('blockquote', 'div')
                  and ('blockquote' in cls or 'cE' in cls or 'cS' in cls)):
                yield 'blockquote', child

    # ── 【重写】通用表格渲染器（保留 colspan/style，剔除 megaEtx） ──────────
    def _render_table(self, table: Tag) -> str:
        rows = []
        for tr in table.find_all('tr'):
            cells = []
            for c in tr.find_all(['td', 'th']):
                # 提取并保留 colspan
                colspan = c.get('colspan', '')
                colspan_attr = f' colspan="{colspan}"' if colspan else ''
                # 提取并保留 style (如 text-align: center)
                style = c.get('style', '')
                style_attr = f' style="{style}"' if style else ''
                
                # 渲染内部内容（会自动走 _inline 处理斜体、公式等）
                content = self._inline(c).strip()
                cells.append(f'<{c.name}{colspan_attr}{style_attr}>{content}</{c.name}>')
            rows.append(f'<tr>{"".join(cells)}</tr>')
        return f'<table>{"".join(rows)}</table>'

    def _render_blockquote(self, elem: Tag) -> str:
        lds   = self._line_divs(elem)
        inner = self._join_lines(lds) if lds else self._inline(elem).strip()
        return f'<blockquote>{inner}</blockquote>'

    def _build_fn_paragraphs(self, elems: List[Tuple[str, Tag]],
                              strip_label: bool) -> List[str]:
        paragraphs: List[str] = []
        cur_parts:  List[str] = []
        first_line = strip_label
        is_first_line_of_para = True
        has_indent = False

        def flush():
            nonlocal is_first_line_of_para, has_indent
            txt = ''.join(cur_parts).strip()
            if txt:
                if has_indent:
                    txt = '\x00INDENT\x00' + txt
                paragraphs.append(txt)
            cur_parts.clear()
            is_first_line_of_para = True
            has_indent = False

        def append_text(text: str):
            nonlocal is_first_line_of_para
            if not text:
                return
            if cur_parts:
                if cur_parts[-1].endswith('-'):
                    cur_parts[-1] = cur_parts[-1][:-1]
                else:
                    cur_parts.append(' ')
            cur_parts.append(text)
            is_first_line_of_para = False

        for etype, elem in elems:
            if etype == 'line':
                lc = elem.find('div', class_=re.compile(r'\blineContent\b'))
                if not lc:
                    continue
                if is_first_line_of_para and 'indent' in (lc.get('class', [])):
                    has_indent = True
                if first_line:
                    clone = BeautifulSoup(str(lc), 'html.parser').find('div')
                    for sup in clone.find_all('sup', class_='megaEtx'):
                        nxt = sup.next_sibling
                        if isinstance(nxt, NavigableString) and re.match(r'\s*\)', str(nxt)):
                            nxt.replace_with('')
                        sup.decompose()
                    fn_span = clone.find('span', class_=re.compile(r'\blabel-footnote\b'))
                    if fn_span:
                        fn_span.decompose()
                    append_text(self._inline(clone).lstrip())
                    first_line = False
                else:
                    append_text(self._inline(lc).strip())
            elif etype == 'table':
                flush()
                paragraphs.append(self._render_table(elem))
            elif etype == 'blockquote':
                flush()
                paragraphs.append(self._render_blockquote(elem))
        flush()
        return paragraphs

    def _parse_note_div(self, note_div: Tag) -> List[RawFootnoteGroup]:
        raw: List[Tuple[Optional[str], List]] = []
        cur_label: Optional[str] = None
        cur_elems: List         = []

        for etype, elem in self._iter_note_elements(note_div):
            if etype == 'line':
                lc = elem.find('div', class_=re.compile(r'\blineContent\b'))
                if lc:
                    span = lc.find('span', class_=re.compile(r'\blabel-footnote\b'))
                    if span:
                        raw.append((cur_label, cur_elems))
                        cur_label = re.sub(r'[^0-9a-z\[\]]+', '',
                                          span.get_text()) or None
                        cur_elems = [(etype, elem)]
                        continue
            cur_elems.append((etype, elem))

        if cur_label or cur_elems:
            raw.append((cur_label, cur_elems))

        return [RawFootnoteGroup(
                    label=label,
                    paragraphs=self._build_fn_paragraphs(
                        elems, strip_label=label is not None))
                for label, elems in raw]

    # ═══════════════════════════════════════════════════════════════════════════
    # 重组层：Block 列表 → HTML
    # ═══════════════════════════════════════════════════════════════════════════

    def _resolve(self, html: str, page_num: int,
                 numberer: FootnoteNumberer) -> str:
        def repl(m):
            kind, old = m.group(1), m.group(2)
            new = numberer.fn_map.get((page_num, old), old)
            if kind == 'FNR':
                anchor_id, href = numberer.next_ref_anchor(new)
                return f'<sup><a id="{anchor_id}" href="{href}">{new})</a></sup>'
            return ''
        result = self._PLACEHOLDER_PAT.sub(repl, html)
        result = re.sub(r'[\s\r\n]{6,}', ' ', result)
        return result

    def _trim_first_page(self, blocks: List[Block]) -> List[Block]:
        for i, blk in enumerate(blocks):
            if ((blk.level in [2,3] and ('Kapitel' in blk.html or "KAPITEL" in blk.html))
                    or (blk.level == 2 and "ABSCHNITT" in blk.html)):
                return blocks[i:]
        return blocks

    def _trim_last_page(self, blocks: List[Block]) -> List[Block]:
        for i, blk in enumerate(blocks):
            if ((blk.level in [2,3] and ('Kapitel' in blk.html or "KAPITEL" in blk.html))
                    or (blk.level == 2 and "ABSCHNITT" in blk.html)):
                valid = list(blocks[:i])
                for blk in blocks[i:]:
                    if blk.type == 'note':
                        valid.append(blk)
                return valid
        return blocks

    # ═══════════════════════════════════════════════════════════════════════════
    # 工具方法
    # ═══════════════════════════════════════════════════════════════════════════

    def _strip_hyphen(self, html: str) -> str:
        m = re.search(r'(-)(\s*(?:</[^>]+>)*)\s*$', html.rstrip())
        if m:
            return html.rstrip()[:m.start(1)] + m.group(2)
        return html.rstrip()

    @staticmethod
    def _splice(base: str, addition: str) -> str:
        b = base.rstrip()
        m = re.search(r'(-)(\s*(?:</[^>]+>)*)\s*$', b)
        if m:
            return b[:m.start(1)] + m.group(2) + addition
        return b + ' ' + addition

    def _check_merge(self, prev: Block, curr: Block) -> Optional[str]:
        if prev is None or curr is None:
            return None
        if prev.type != curr.type:
            return None
        prev_text = re.sub(r'<[^>]+>', '', prev.html).strip()
        if not prev_text:
            return None
        last_char = prev_text[-1]
        if last_char == '-':
            return 'hyphen'
        if last_char.isalnum() or last_char == ',':
            if curr.indent:
                return None
            return 'word/comma'
        if not curr.indent and prev.indent:
            return 'word/comma'
        return None

    # ═══════════════════════════════════════════════════════════════════════════
    # 重组核心：正文先行 → 提取编号 → 按需取注
    # ═══════════════════════════════════════════════════════════════════════════

    def merge_group(self, pages: List[int], title: str = '',startpage:int=0,endpage:int=0) -> str:
        numberer = FootnoteNumberer()
        fn_dict: Dict[str, Footnote] = {}
        parts:   List[str]           = []

        pending: Optional[Block]    = None
        open_fn: Optional[Footnote] = None

        def flush():
            nonlocal pending
            if pending is None:
                return
            blk = pending
            cls = ' class="ni"' if not blk.indent else ''
            if blk.type == 'p':
                parts.append(f'<p{cls}>{blk.html}</p>')
            elif blk.type == 'blockquote':
                parts.append(f'<blockquote{cls}>{blk.html}</blockquote>')
            pending = None

        def flush_page_num(page_num):
            nonlocal pending
            if pending is None:
                return
            blk = pending
            cls = ' class="ni"' if not blk.indent else ''
            if blk.type == 'p':
                parts.append(f'<p{cls}>{blk.html}</p>')
                parts.append(f'<a id="S{page_num}"></a>')
            elif blk.type == 'blockquote':
                parts.append(f'<blockquote{cls}>{blk.html}</blockquote>')
                parts.append(f'<a id="S{page_num}"></a>')
            pending = None

        for pg_idx, page_num in enumerate(pages):
            path = self._path(page_num)
            if not path.exists():
                anchor = f'<a id="S{page_num}"></a>'
                if pending:
                    pending.html += anchor
                else:
                    parts.append(anchor)
                continue

            soup   = BeautifulSoup(path.read_text(encoding='utf-8'), 'html.parser')
            blocks = self._parse_page(soup)

            if self.cutcontent:
                if pg_idx == 0:
                    blocks = self._trim_first_page(blocks)
                if pg_idx == len(pages) - 1:
                    blocks = self._trim_last_page(blocks)

            # ── 脚注收集 ──
            if self.use_footnotes:
                for blk in blocks:
                    if blk.type != 'note':
                        continue
                    for rfg in blk.footnotes:
                        can_merge = False
                        if not rfg.label and not open_fn:
                            continue
                        if open_fn:
                            if rfg.label is None or rfg.label == open_fn.old_num:
                                can_merge = True

                        if can_merge and rfg.paragraphs:
                            target_html = (open_fn.extra_paras[-1]
                                           if open_fn.extra_paras
                                           else open_fn.content_html)
                            next_html = rfg.paragraphs[0]
                            is_indent = next_html.startswith('\x00INDENT\x00')
                            if is_indent:
                                next_html = next_html[len('\x00INDENT\x00'):]
                            prev_text = re.sub(r'<[^>]+>', '', target_html).strip()
                            is_ended  = prev_text.endswith(('.', '!', '?', '…'))
                            if not is_ended and not is_indent:
                                spliced = self._splice(target_html, next_html)
                                if open_fn.extra_paras:
                                    open_fn.extra_paras[-1] = spliced
                                else:
                                    open_fn.content_html = spliced
                                open_fn.extra_paras.extend(rfg.paragraphs[1:])
                            else:
                                open_fn = None
                        else:
                            if rfg.label is not None:
                                key = (page_num, rfg.label)
                                if key not in numberer.fn_map:
                                    new_num     = numberer.assign(page_num, rfg.label)
                                    clean_paras = [
                                        p.replace('\x00INDENT\x00', '')
                                        for p in rfg.paragraphs
                                    ]
                                    fn = Footnote(
                                        rfg.label, new_num, page_num,
                                        clean_paras[0] if clean_paras else '',
                                        list(clean_paras[1:]))
                                    fn_dict[fn.new_num] = fn
                                    open_fn = fn

            # ── 正文重组 ──
            first_content = True
            pre_hn_lv=9
            heading_part=""
            for blk in blocks:
                if blk.type == 'note':
                    continue

                html = self._resolve(blk.html, page_num, numberer)

                # --- 标题与编者按 ---
                if (blk.type.startswith('h') and blk.level>0) or blk.type == 'editorial':
                    flush()
                    if blk.type == 'editorial':
                        parts.append(f'<div class="ed">{html}</div>')
                    else:
                        hn    = f'h{min(blk.level, 6)}'
                        extra = f' class="h{blk.level}"' if blk.level > 6 else ''
                        head=f'<{hn}{extra}>'
                        if not first_content:
                            if blk.level==pre_hn_lv:
                                heading_part+="<br>"+html
                                continue
                                #parts.append('<br>')
                            elif blk.level!=pre_hn_lv and heading_part:
                                parts.append(heading_part+f'</h{min(pre_hn_lv, 6)}>')
                                pre_hn_lv=min(blk.level, 6)
                        heading_part=head+html

                                #parts.append(f'</h{min(pre_hn_lv, 6)}>')
                        #if first_content or not heading_part:
                            #parts.append(head+html)
                        #    heading_part+=head+html
                        #else:
                            #parts.append(html)
                    first_content = False
                    pre_hn_lv=min(blk.level, 6)
                    continue
                if pre_hn_lv>0 and pre_hn_lv<9 and blk.level==0 and heading_part:
                    parts.append(heading_part+f'</h{min(pre_hn_lv, 6)}>')
                    pre_hn_lv=9
                    heading_part=""
                

                # --- 表格 ---
                if blk.type == 'table':
                    if first_content and pg_idx > 0:
                        flush_page_num(page_num)
                    else:
                        flush()
                    resolved_html = self._resolve(html, page_num, numberer)
                    parts.append(resolved_html)
                    first_content = False
                    continue

                # --- 段落 ---
                curr_blk = Block(type=blk.type, html=html,
                                 indent=blk.indent, quote=blk.quote)

                if first_content and pg_idx > 0:
                    anchor = f'<a id="S{page_num}"></a>'
                    if pending:
                        merge_type = self._check_merge(pending, curr_blk)
                        if merge_type == 'hyphen':
                            clean_prev       = self._strip_hyphen(pending.html)
                            curr_blk.html    = clean_prev+ "fanchorf" + curr_blk.html 
                            curr_blk.html=re.sub(r""" ([\S]+?)fanchorf([\S]+?) """,r" \1\2"+anchor+r" ",curr_blk.html,flags=re.IGNORECASE|re.DOTALL)
                            curr_blk.indent  = pending.indent
                            pending          = curr_blk
                        elif merge_type == 'word/comma':
                            curr_blk.html   = pending.html.rstrip() + ' ' + anchor + curr_blk.html
                            curr_blk.indent = pending.indent
                            pending         = curr_blk
                        else:
                            flush_page_num(page_num)
                            pending = curr_blk
                    else:
                        pending = curr_blk
                else:
                    if pending:
                        flush()
                    pending = curr_blk

                first_content = False

        flush()

        body_text     = '\n'.join(parts)
        seen:          set       = set()
        ordered_nums: List[str] = []
        for m in re.finditer(r'id="ZM([\da-z\[\]]+-?\d*)"', body_text):
            raw = m.group(1)
            num = re.sub(r'-\d+$', '', raw)
            if num not in seen:
                seen.add(num)
                ordered_nums.append(num)
        fn_html=""  
        if fn_dict:
            fn_html = self._fn_html_by_order(ordered_nums, fn_dict, numberer)
        vol=self.volume
        if self.volume==1:
            if startpage>=311:
                vol="1.2"
            else:
                vol="1.1"
        if startpage==endpage:
            
            que=f"""<div class="que">Quelle: MEGA II/{vol}: {self.volume_info}, Seite {startpage}</div>"""
        else:
            que=f"""<div class="que">Quelle: MEGA II/{vol}: {self.volume_info}, Seite {startpage}-{endpage}</div>"""

        t = title or f"II.{self.volume}"
        return (f'<!DOCTYPE html>\n<html lang="de">\n<head>\n'
                f'<meta charset="UTF-8">\n<title>{t}</title>\n'
                f'<link rel="stylesheet" href="{CSS_HREF}">\n'
                f'</head>\n<body>\n{body_text}\n{fn_html}\n{que}</body>\n</html>')

    @staticmethod
    def _fn_html_by_order(ordered_nums: List[str],
                          fn_dict: Dict,
                          numberer: FootnoteNumberer) -> str:
        if not ordered_nums:
            return ''
        lines = ['<aside class="fn">', '<div class="fnt">Fußnoten</div>']
        for num in ordered_nums:
            fn = fn_dict.get(num)
            if fn is None:
                continue
            content  = fn.content_html.replace('\x00INDENT\x00', '')

            ref_aids = numberer.ref_anchors_for(fn.new_num)
            num_link = (f'<a id="M{fn.new_num}" href="#{ref_aids[0]}">'
                        f'{fn.new_num})</a> ')
            extra_back = ''.join(
                f'<a href="#{aid}">↑{i + 2}</a> '
                for i, aid in enumerate(ref_aids[1:])
            )
            if "<blockquote" in content or "<table" in content:
                lines.append(f'<p class="fni">{num_link}{extra_back}</p><div>{content}</div>')
            else:
                lines.append(f'<p class="fni">{num_link}{extra_back} {content}</p>')
            for ep in fn.extra_paras:
                ep_clean = ep.replace('\x00INDENT\x00', '')
                wrap = ('div'
                        if '<table' in ep_clean or '<blockquote' in ep_clean
                        else 'p')
                lines.append(f'<{wrap}>{ep_clean}</{wrap}>')
        lines.append('</aside>')
        return '\n'.join(lines)

    # ── entry point ───────────────────────────────────────────────────────────

    def run(self, titles: Dict[int, str] = None) -> None:
        self.out_dir.mkdir(parents=True, exist_ok=True)
        titles = titles or {}
        index_content=""

        for group in self.groups:
            start = group[0]
            print(f"  p{start:04d}–p{group[-1]:04d}  ({len(group)} pages)")
            merged = self.merge_group(group, titles.get(start, ''),start,group[-1])
            out    = self.out_dir / f"MEGA2-II-{self.volume:02d}-{start:04d}.html"
            out.write_text(merged, encoding='utf-8', newline='\n')
            print(f"    → {out.name}")
            index_content+=f"""<a href="MEGA2-II-{self.volume:02d}-{start:04d}.html">{start}</a><br>\n"""
        index_file=self.out_dir /"index.html"
        index_file.write_text(index_content,encoding="utf-8")


def main():
    volumes = [1,5, 10, 13, 15]
    excel_file=Path(r"LENIN-toc.xlsx")
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb.active
    ws = wb['Sheet1']  

    for volume in volumes:
        for row in ws.iter_rows(min_row=volume, max_row=volume, values_only=True):
            volume_info=row[6]
        page_group = MEGAok.page_group[volume]
        input_dir  = Path(f'./mega_raw/II_{volume}/raw')
        output_dir = Path(f'./MEGA_II_pre/{volume}')
        merger = MegaEtxMerger(volume=volume, input_dir=input_dir,
                               output_dir=output_dir, page_group=page_group,volume_info=volume_info)
        merger.run()


if __name__ == "__main__":
    main()